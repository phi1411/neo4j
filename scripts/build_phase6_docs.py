"""Build the two Vietnamese submission documents from verified project evidence.

Run with a documentation Python environment containing python-docx and Pillow.
Application dependencies and database contents are unaffected.
"""
from pathlib import Path
import json
import re
import textwrap

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
DATA = ROOT / 'database'
SHAPES = json.loads((DATA / 'shapes.json').read_text(encoding='utf-8'))
PROPS = json.loads((DATA / 'properties.json').read_text(encoding='utf-8'))
CONDS = json.loads((DATA / 'conditions.json').read_text(encoding='utf-8'))
P2 = json.loads((DOCS / 'PHASE_2_RESULTS.json').read_text(encoding='utf-8'))
P5 = json.loads((DOCS / 'PHASE_5_RESULTS.json').read_text(encoding='utf-8'))
assert P2['status'] == P5['status'] == 'PASS'
assert (P2['nodes'], P2['relationships']) == (44, 79)
assert (len(SHAPES), len(PROPS), len(CONDS)) == (7, 19, 18)
NAME = {s['id']: s['name'] for s in SHAPES}
PNAME = {p['id']: p['name'] for p in PROPS}
QUERY_TEXT = (DATA / 'demo_queries.cypher').read_text(encoding='utf-8')
QUERIES = {}
for part in re.split(r'^// @query ', QUERY_TEXT, flags=re.M)[1:]:
    key, body = part.split('\n', 1)
    QUERIES[key.strip()] = '\n'.join(line for line in body.strip().splitlines() if not line.startswith('//')).strip()
assert len(QUERIES) == 15
QUERY_INFO = [
    ('Q01', 'Toàn bộ graph', '44 node, 79 relationship'),
    ('Q02', 'Tính chất hiệu lực của HV', '12 tính chất, kèm nơi khai báo'),
    ('Q03', 'Cha trực tiếp của HV', 'HCN, HTHOI'),
    ('Q04', 'Mọi tổ tiên của HV', 'HBH, HCN, HTHOI, TQ'),
    ('Q05', 'Trường hợp đặc biệt của HBH', 'HCN, HTHOI, HV'),
    ('Q06', 'Bốn cạnh bằng nhau', 'HTHOI, HV'),
    ('Q07', 'Bốn góc vuông', 'HCN, HV'),
    ('Q08', 'Hai đường chéo bằng nhau', 'HCN, HTC, HV'),
    ('Q09', 'Đường đi TQ – HV', '2 đường đi, mỗi đường dài 3 cạnh'),
    ('Q10', 'Tính chất chung HV và HCN', '9 tính chất'),
    ('Q11', 'Tính chất chung HV và HTHOI', '10 tính chất'),
    ('Q12', 'Đường chéo vuông góc và chia đôi nhau', 'HTHOI, HV'),
    ('Q13', 'Bậc và quan hệ từng hình', '7 dòng, phân biệt hướng vào/ra'),
    ('Q14', 'Lân cận HV độ sâu 1–3', '8/11; 23/43; 39/71 (node/cạnh)'),
    ('Q15', 'Giải thích tính chất kế thừa', '2 đường HV → HCN/HTHOI → HBH → Property'),
]
IMAGES = DOCS / 'images'
FIVE = IMAGES / 'phase5'
TOPIC = 'MÔ HÌNH HOÁ – TỨ GIÁC TRONG HÌNH HỌC PHẲNG VỚI NEO4J'
PARAM = ":params {dataset:'quadrilateral-v1',shape_id:'HV',depth:3}"
PY = r'.\.venv\Scripts\python.exe'


def field(p, instruction):
    run = p.add_run()
    start = OxmlElement('w:fldChar'); start.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([start, text, end])


def new_doc(subtitle):
    d = Document()
    # Remove theme fonts and decorative paragraph borders from the default template.
    for style in d.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'
            fonts = style._element.get_or_add_rPr().rFonts
            for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
                fonts.attrib.pop(qn('w:' + attr), None)
            fonts.set(qn('w:cs'), 'Times New Roman')
            if style._element.pPr is not None:
                for border in list(style._element.pPr.findall(qn('w:pBdr'))):
                    style._element.pPr.remove(border)
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    section.header_distance, section.footer_distance = Cm(0.9), Cm(0.9)
    section.different_first_page_header_footer = True
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        style = d.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal = d.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for level, size in [(1, 16), (2, 13), (3, 12)]:
        s = d.styles[f'Heading {level}']; s.font.size = Pt(size); s.font.bold = True
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    d.styles['Caption'].font.size = Pt(10.5)
    d.styles['Caption'].font.italic = True
    header = section.header.paragraphs[0]
    header.text = 'NoSQL • Tứ giác và Neo4j'; header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Trang ').font.size = Pt(10); field(footer, 'PAGE')
    settings = d.settings.element
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
    for value in ['[Trường]', '[Khoa]', 'MÔN HỌC: NoSQL']:
        p = d.add_paragraph(value); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
    p = d.add_paragraph(subtitle, 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48); p.runs[0].font.size = Pt(20); p.runs[0].bold = True
    p = d.add_paragraph(TOPIC); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(17); p.runs[0].bold = True; p.paragraph_format.space_after = Pt(48)
    for value in ['Sinh viên: [Họ tên sinh viên]', 'MSSV: [MSSV]', 'Lớp: [Lớp]', 'Giảng viên: [Giảng viên]', 'Năm học: [Năm học]']:
        p = d.add_paragraph(value); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph('Tài liệu dựa trên phiên bản đã nghiệm thu Phase 5, tháng 10/2026.')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(32)
    h(d, 'MỤC LỤC', page=True, toc=False)
    p = d.add_paragraph(); field(p, 'TOC \\o "1-2" \\h \\z \\u')
    d.core_properties.title = subtitle + ' – ' + TOPIC
    d.core_properties.subject = 'Bài tập môn NoSQL; dataset quadrilateral-v1'
    d.core_properties.author = '[Họ tên sinh viên]'
    d.core_properties.keywords = 'Neo4j, Cypher, tứ giác, NoSQL'
    return d


def h(d, text, level=1, page=False, toc=True):
    p = d.add_paragraph(text, f'Heading {level}' if toc else 'Title')
    if page: p.paragraph_format.page_break_before = True
    return p


def para(d, text):
    return d.add_paragraph(text)


def code(d, value):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    lines = []
    for line in value.splitlines():
        indent = ' ' * (len(line) - len(line.lstrip()))
        lines.extend(textwrap.wrap(line, 88, subsequent_indent=indent + '  ', break_long_words=False, break_on_hyphens=False, replace_whitespace=False) or [''])
    run = p.add_run('\n'.join(lines)); run.font.name = 'Consolas'; run.font.size = Pt(8.5)
    return p


def table(d, caption, headers, rows, widths, compact=False):
    p = d.add_paragraph(caption, 'Caption'); p.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    for c, width in zip(t.columns, widths): c.width = Cm(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Cm(width); cell.text = label
        sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E7E6E6'); cell._tc.get_or_add_tcPr().append(sh)
        for run in cell.paragraphs[0].runs: run.bold = True
    rep = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells = t.add_row().cells
        for cell, value, width in zip(cells, row, widths): cell.width = Cm(width); cell.text = str(value)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + side); node.set(qn('w:w'), '25' if compact and side in ['top','bottom'] else '80'); node.set(qn('w:type'), 'dxa'); margins.append(node)
            pr.append(margins)
            valign = OxmlElement('w:vAlign'); valign.set(qn('w:val'), 'center'); pr.append(valign)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:' + side); border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '4'); border.set(qn('w:color'), 'D9D9D9'); borders.append(border)
            pr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.05
                for run in p.runs: run.font.size = Pt(10 if compact else 10.5)
    return t


def figure(d, filename, caption, width=16, max_height=15.5):
    path = filename if isinstance(filename, Path) else FIVE / filename
    with Image.open(path) as im: w, height = im.size
    width = min(width, max_height * w / height)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    p = d.add_paragraph(caption, 'Caption'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def overview():
    d = new_doc('BÁO CÁO TỔNG QUAN DỰ ÁN')
    h(d, 'CHƯƠNG 1 — GIỚI THIỆU', page=True)
    for title, text in [
        ('1.1 Lý do chọn đề tài', 'Kiến thức về tứ giác thường được trình bày thành các định nghĩa, tính chất và dấu hiệu nhận biết riêng lẻ. Khi xét hình vuông, người học cần liên hệ đồng thời với hình chữ nhật, hình thoi và hình bình hành. Đề tài dùng một đồ thị tri thức nhỏ để lưu các liên hệ này thành dữ liệu có thể truy vấn, quan sát và giải thích.'),
        ('1.2 Mục tiêu', 'Mục tiêu chính là mô hình hoá bảy loại tứ giác bằng Neo4j, truy vấn bằng Cypher và xây dựng giao diện tiếng Việt phục vụ demo môn NoSQL. Hệ thống phải tìm được tính chất kế thừa, lớp cha, lớp tổ tiên, các tính chất chung và đường đi giải thích nguồn gốc một tính chất.'),
        ('1.3 Phạm vi', 'Chỉ xét tứ giác lồi, không suy biến trong hình học phẳng. Dataset gồm Tứ giác, Hình thang, Hình thang cân, Hình bình hành, Hình chữ nhật, Hình thoi và Hình vuông. Ứng dụng lưu kiến thức về loại hình, không lưu các hình có toạ độ và không tự đo cạnh hoặc góc từ ảnh.'),
        ('1.4 Kết quả cần đạt', 'Bài nộp gồm báo cáo tổng quan, Data + Source App và hướng dẫn sử dụng có ảnh kết quả thật. Phiên bản đã nghiệm thu có database quadrilateral, dataset quadrilateral-v1, backend Flask, sáu trang web, 15 truy vấn demo và các báo cáo kiểm thử. Việc công bố link GitHub sẽ thực hiện sau kiểm tra secret ở Phase 7.')]:
        h(d, title, 2); para(d, text)
    h(d, 'CHƯƠNG 2 — CƠ SỞ LÝ THUYẾT', page=True)
    for title, text in [
        ('2.1 NoSQL', 'NoSQL bao gồm nhiều mô hình lưu trữ như document, key-value, wide-column và graph. Trong đề tài này, lựa chọn mô hình graph xuất phát từ nhu cầu đi qua các quan hệ phân loại và chia sẻ tính chất; không xuất phát từ kích thước dữ liệu lớn.'),
        ('2.2 Graph Database', 'Graph database biểu diễn các thực thể bằng node và các liên hệ bằng relationship. Một đường đi có thể chứa nhiều relationship liên tiếp, giúp truy vấn lớp tổ tiên hoặc lân cận mà không cần cố định số cấp trong dữ liệu.'),
        ('2.3 Neo4j', 'Neo4j sử dụng property graph: node và relationship có thể mang thuộc tính; node có label, relationship có type và hướng. Project đã chạy trên Neo4j Enterprise 2026.09.0. Các thao tác đọc dùng Neo4j Python Driver; các script Cypher đảm nhiệm schema, seed và kiểm tra dữ liệu.'),
        ('2.4 Node', 'Node đại diện cho một loại hình, một mệnh đề tính chất hoặc một điều kiện nhận biết. Ví dụ Shape có id HV, Property có id FOUR_EQUAL_SIDES và Condition có id C_HV_05.'),
        ('2.5 Relationship', 'Relationship thể hiện ý nghĩa cụ thể giữa hai node. IS_A đi từ loại đặc biệt tới loại tổng quát; HAS_PROPERTY đi từ loại hình tới tính chất được khai báo trực tiếp. Hướng của cạnh là một phần của mô hình.'),
        ('2.6 Property', 'Thuộc tính lưu trên node như id, name, definition, dataset là các cặp khoá–giá trị. Cần phân biệt cơ chế property của Neo4j với label Property trong project: label này đại diện cho mệnh đề hình học được chia sẻ giữa nhiều hình và điều kiện.'),
        ('2.7 Cypher', 'Cypher là ngôn ngữ truy vấn theo mẫu đồ thị. MATCH tìm mẫu; WHERE giới hạn dữ liệu; RETURN trả kết quả. Pattern có độ dài biến thiên dùng để duyệt kế thừa. DISTINCT loại kết quả lặp khi một tính chất được tới qua hai nhánh của hình vuông.')]:
        h(d, title, 2); para(d, text)
    h(d, 'CHƯƠNG 3 — PHÂN TÍCH BÀI TOÁN', page=True)
    h(d, '3.1 Các loại tứ giác', 2)
    table(d, 'Bảng 1. Bảy Shape trong dataset', ['ID', 'Loại hình', 'Định nghĩa'], [(s['id'], s['name'], s['definition']) for s in SHAPES], [1.5, 3.5, 11.5])
    h(d, '3.2 Quy ước hình thang', 2)
    para(d, 'Project quy ước hình thang có đúng một cặp cạnh đối song song. Vì vậy, nhánh hình thang và nhánh hình bình hành tách riêng dưới Tứ giác. Không có cạnh HBH IS_A HT hoặc HCN IS_A HTC. Đây là quy ước taxonomy của project; các tài liệu dùng định nghĩa hình thang có ít nhất một cặp cạnh song song có thể có hệ phân loại khác.')
    h(d, '3.3 Tính chất', 2, page=True)
    para(d, 'Danh mục gồm 19 mệnh đề, chia thành cấu trúc, cạnh, góc, đường chéo và đường chéo–góc. Ba mệnh đề THREE_RIGHT_ANGLES, HAS_RIGHT_ANGLE và ADJACENT_SIDES_EQUAL được dùng làm giả thiết của dấu hiệu nhận biết; không có cạnh HAS_PROPERTY khai báo trực tiếp cho các mệnh đề này. Danh mục thể hiện tri thức đã chọn, không phải tập tất cả hệ quả toán học có thể suy ra.')
    table(d, 'Bảng 2. Danh mục Property', ['Mệnh đề hình học', 'Nhóm'], [(p['name'], {'structure':'Cấu trúc','side':'Cạnh','angle':'Góc','diagonal':'Đường chéo','diagonal_angle':'Đường chéo–góc'}[p['category']]) for p in PROPS], [12, 4.5], compact=True)
    h(d, '3.4 Dấu hiệu nhận biết', 2)
    para(d, 'Một dấu hiệu nhận biết phải có bối cảnh và các giả thiết đủ. Chẳng hạn, hai đường chéo bằng nhau riêng lẻ không đủ kết luận hình chữ nhật: hình thang cân cũng có tính chất này. Condition C_HCN_03 yêu cầu ngữ cảnh Hình bình hành cùng Property Hai đường chéo bằng nhau mới nhận biết Hình chữ nhật.')
    h(d, '3.5 Vì sao bài toán phù hợp Graph Database', 2)
    para(d, 'Hình vuông có hai lớp cha và nhiều đường đi tới cùng một tổ tiên. Property là thực thể dùng chung, còn Condition nối loại cần nhận biết với các tiền đề. Truy vấn Q15 trả cả đường đi giải thích vì sao hình vuông có đường chéo chia đôi nhau. SQL vẫn mô hình hoá được bằng bảng liên kết và truy vấn đệ quy; đề tài không có benchmark để kết luận Neo4j nhanh hơn SQL.')
    h(d, 'CHƯƠNG 4 — THIẾT KẾ GRAPH', page=True)
    h(d, '4.1 Node Shape', 2)
    para(d, 'Shape lưu id ổn định, tên tiếng Việt, định nghĩa, tên đã chuẩn hoá để tìm kiếm, thứ tự hiển thị và nguồn tham khảo. Mọi node thuộc project có thêm label ProjectEntity và thuộc tính dataset = quadrilateral-v1. Khoá duy nhất là cặp (dataset, id), tránh xung đột khi một database chứa dataset khác.')
    h(d, '4.2 Node Property', 2)
    para(d, 'Property lưu id, name, category, description, notation và source_ref. Một mệnh đề như DIAGONALS_EQUAL chỉ có một node, được dùng bởi Hình chữ nhật, Hình thang cân và các Condition phù hợp. Việc tái sử dụng node giúp truy vấn tính chất chung có ý nghĩa dữ liệu.')
    h(d, '4.3 Node Condition', 2)
    para(d, 'Condition lưu id, name, statement, explanation, logic = AND và source_ref. Các tiền đề nằm trên cạnh REQUIRES_SHAPE và REQUIRES_PROPERTY; Shape đích nối tới Condition bằng HAS_CONDITION. Tứ giác nền không cần dấu hiệu nhận biết riêng trong dataset. Tổng số Condition là 18.')
    h(d, '4.4 Relationships', 2)
    table(d, 'Bảng 3. Ý nghĩa và số lượng relationship', ['Type', 'Hướng và ý nghĩa', 'Số'], [
        ('IS_A','Shape → Shape: là trường hợp đặc biệt của',7),
        ('HAS_PROPERTY','Shape → Property: khai báo trực tiếp',17),
        ('HAS_CONDITION','Shape → Condition: dấu hiệu nhận biết',18),
        ('REQUIRES_SHAPE','Condition → Shape: ngữ cảnh cần có',18),
        ('REQUIRES_PROPERTY','Condition → Property: giả thiết cần có',19)], [4.5,10.5,1.5])
    h(d, '4.5 IS_A taxonomy', 2, page=True)
    para(d, 'Hình thang cân → Hình thang → Tứ giác. Hình chữ nhật và Hình thoi → Hình bình hành → Tứ giác. Hình vuông → Hình chữ nhật và Hình thoi. Mũi tên trong sơ đồ đi từ lớp đặc biệt tới lớp tổng quát; đồ thị phân loại có bảy cạnh, không có chu trình.')
    figure(d, '02_taxonomy_graph.jpg', 'Hình 1. Taxonomy thật trên ứng dụng, thể hiện hai lớp cha của Hình vuông.')
    h(d, '4.6 Kế thừa Property', 2)
    para(d, 'Tính chất hiệu lực là hợp của HAS_PROPERTY tại chính hình và tại mọi tổ tiên qua IS_A. HV không lưu HAS_PROPERTY trực tiếp nhưng có 12 tính chất hiệu lực. Q02 giữ thông tin nơi khai báo; DISTINCT loại lặp do đa kế thừa. Không sao chép cạnh tính chất từ lớp cha xuống lớp con, và không kế thừa Condition.')
    h(d, '4.7 Condition AND/OR', 2, page=True)
    para(d, 'Trong một Condition, ngữ cảnh và tất cả Property phải đồng thời thoả mãn (AND). Nhiều Condition nối với cùng Shape là các cách nhận biết thay thế (OR). Ứng dụng hiển thị cấu trúc và giải thích điều kiện; chưa có bộ suy luận tự động nhận đầu vào rồi kết luận một loại hình.')
    table(d, 'Bảng 4. Năm cách nhận biết Hình vuông (OR giữa các dòng)', ['Condition', 'Ngữ cảnh', 'Giả thiết đồng thời'], [(c['id'],NAME[c['context']], '\n'.join(PNAME[p] for p in c['properties'])) for c in CONDS if c['target']=='HV'], [2.8,3.5,10.2])
    para(d, 'Ví dụ C_HV_05: Tứ giác AND Có bốn cạnh bằng nhau AND Có bốn góc vuông ⇒ Hình vuông. C_HV_02: Hình chữ nhật AND Hai đường chéo vuông góc ⇒ Hình vuông. Một Property xuất hiện trong Condition không đồng nghĩa là đủ để nhận biết khi thiếu ngữ cảnh.')
    h(d, '4.8 Số lượng dữ liệu', 2)
    table(d, 'Bảng 5. Thống kê dataset đã kiểm thử', ['Thành phần', 'Số lượng'], [('Shape',7),('Property',19),('Condition',18),('Tổng node',44),('Tổng relationship',79),('Unique constraint',3),('Index ONLINE, gồm index hỗ trợ constraint',4)], [12,4.5])
    para(d, 'Seed dùng MERGE theo ID ổn định. Phase 5 seed lại hai lần đều không tạo thêm node hoặc relationship; snapshot của dataset và dữ liệu ngoài project giữ nguyên. Script không xóa database hay node ngoài phạm vi dataset.')
    h(d, 'CHƯƠNG 5 — XÂY DỰNG HỆ THỐNG', page=True)
    for title,text in [
        ('5.1 Kiến trúc', 'Trình duyệt → Flask → Neo4j Python Driver → Neo4j. HTML/CSS/JavaScript trình bày dữ liệu và vis-network vẽ đồ thị. Neo4j là database chính; website không dùng dữ liệu mẫu thay thế khi kết nối lỗi.'),
        ('5.2 Neo4j', 'Database quadrilateral lưu dataset quadrilateral-v1. constraints.cypher và indexes.cypher tạo schema; seed.cypher tạo dữ liệu; demo_queries.cypher lưu 15 mẫu truy vấn; validation.cypher kiểm tra tính toàn vẹn. Các catalog JSON là nguồn dữ liệu để sinh seed có thể tái lập.'),
        ('5.3 Flask', 'Backend dùng app factory, cấu hình từ biến môi trường hoặc .env và một driver Neo4j dùng chung theo vòng đời ứng dụng. Service thực hiện các truy vấn đọc có tham số, timeout và chuyển lỗi kết nối thành thông báo phù hợp. run.py chạy server tại 127.0.0.1:5000, tắt debug và reloader.'),
        ('5.4 API', 'API cung cấp thống kê, danh sách/chi tiết hình, tìm kiếm, taxonomy, graph lân cận và truy vấn demo. Payload thành công có success = true và data; lỗi có success = false và error. Mã HTTP gồm 400 đầu vào sai, 404 không tìm thấy, 500 lỗi truy vấn/nội bộ và 503 lỗi cấu hình hoặc kết nối. GET/POST /api/queries/<id> chỉ chạy 15 truy vấn trong whitelist.'),
        ('5.5 Frontend', 'Giao diện tiếng Việt có Navbar và sáu trang: Tổng quan, Tứ giác, Chi tiết, Graph, Truy vấn và Giới thiệu. Các trang dùng kết quả API để hiển thị. Khi dữ liệu lỗi, ứng dụng xoá kết quả cũ và báo lỗi, tránh trình bày dữ liệu cũ như kết quả truy vấn mới.'),
        ('5.6 vis-network', 'Thư viện vis-network được lưu trong static của project. Node Shape, Property và Condition có hình dạng/màu khác nhau; cạnh có hướng và nhãn. Người dùng có thể kéo, thu phóng, căn toàn bộ graph, chọn node/cạnh và đọc thông tin liên quan.'),
        ('5.7 Search', 'Chuỗi tìm kiếm được chuẩn hoá chữ hoa/thường, dấu tiếng Việt và đ/đ dạng không dấu. Nhập hinh vuong vẫn trả Hình vuông. Truy vấn dùng tham số; API giới hạn độ dài và số kết quả. Ô tìm kiếm trống trên trang danh sách hiển thị cả bảy hình; API search từ chối chuỗi trống bằng HTTP 400.'),
        ('5.8 Graph visualization', 'Taxonomy chỉ hiển thị bảy Shape và bảy cạnh IS_A. Graph lân cận duyệt các loại relationship theo cả hai hướng với độ sâu 1–3 và trả các cạnh giữa node được chọn. Với HV: độ sâu 1 là 8 node/11 cạnh, độ sâu 2 là 23/43, độ sâu 3 là 39/71. Lân cận bao gồm điều kiện, tiền đề và tính chất, nên không phải chỉ là các tổ tiên của HV.')]:
        h(d,title,2,page=title.startswith('5.5')); para(d,text)
    h(d, 'CHƯƠNG 6 — TRUY VẤN CYPHER', page=True)
    para(d, 'Các kết quả sau lấy từ kiểm thử thật Phase 2 và được đối chiếu lại ở Phase 5. Chạy trong database quadrilateral. Neo4j Query/Browser cần khai báo tham số; dấu :params là lệnh của Browser, không phải cú pháp Cypher gửi qua Python Driver.')
    code(d, PARAM)
    table(d, 'Bảng 6. Tổng hợp 15 truy vấn demo', ['Mã','Mục đích','Kết quả đã kiểm thử'], QUERY_INFO, [1.3,6.2,9])
    meanings = {
        'Q02':'Truy vấn hợp các tính chất trên đường kế thừa và giữ nơi khai báo. Hình vuông nhận hai đường chéo bằng nhau từ HCN, đường chéo chia đôi nhau từ HBH; DISTINCT xử lý các đường trùng qua hai nhánh.',
        'Q04':'Pattern có độ dài biến thiên thể hiện traversal tới các tổ tiên. HT và HTC không nằm trong kết quả, phù hợp taxonomy đã chọn.',
        'Q06':'Hình vuông vẫn được tìm thấy dù không có HAS_PROPERTY trực tiếp. Vì vậy truy vấn không chỉ kiểm tra cạnh tại node đang xét.',
        'Q08':'Kết quả có cả Hình thang cân và Hình chữ nhật. Đây là ví dụ cho thấy một tính chất riêng lẻ không phải điều kiện đủ để phân loại hình.',
        'Q09':'Pattern liệt kê các đường IS_A có độ dài 1–6, không dùng shortestPath. Trong taxonomy hiện tại, hai đường trả về đều dài ba cạnh và cũng là ngắn nhất. Pattern đi không định hướng; hướng lưu trữ IS_A vẫn là con → cha.',
        'Q10':'Giao hai tập tính chất hiệu lực cho thấy chín tính chất của Hình chữ nhật đều có ở Hình vuông, không cần nhân bản node Property.',
        'Q12':'Hai yêu cầu về đường chéo được kiểm tra đồng thời. Các đường kế thừa có thể khác nhau, nhưng đều thuộc cùng Shape cần tìm.',
        'Q15':'Kết quả trả đường đi để giải thích tính chất: HV → HCN → HBH → DIAGONALS_BISECT_EACH_OTHER và HV → HTHOI → HBH → cùng Property. Đây là minh chứng trực tiếp cho giá trị traversal của graph.'}
    info = {q:(purpose,result) for q,purpose,result in QUERY_INFO}
    for i,q in enumerate(meanings,1):
        h(d, f'6.{i} {q} — {info[q][0]}', 2, page=q in ['Q02','Q06','Q09','Q12'])
        para(d,'Mục tiêu: '+ info[q][0] + '.')
        code(d,QUERIES[q])
        para(d,'Kết quả thực tế: '+ info[q][1] + '.')
        para(d,'Ý nghĩa: '+ meanings[q])
    h(d,'CHƯƠNG 7 — KẾT QUẢ',page=True)
    para(d,'Database kết nối thật và có 44 node, 79 relationship. Backend phục vụ dữ liệu cho dashboard, chi tiết và graph. Website tìm kiếm được cả tiếng Việt có dấu và không dấu. Các hình dưới đây là ảnh chụp ứng dụng trong kiểm thử, không phải ảnh mô phỏng.')
    figure(d,'01_dashboard.jpg','Hình 2. Dashboard thực tế: thống kê dataset và graph tổng quan.')
    para(d,'Dashboard làm điểm bắt đầu demo: người xem thấy số lượng dữ liệu và có thể mở chi tiết Hình vuông hoặc trang Truy vấn. Các số thống kê lấy từ Neo4j, không được gán cố định trong giao diện.')
    h(d,'7.1 Chi tiết và tính chất kế thừa',2,page=True)
    figure(d,'05_shape_hv_properties.jpg','Hình 3. Trang Hình vuông hiển thị 12 tính chất hiệu lực cùng nguồn khai báo.')
    para(d,'Danh sách tính chất cho thấy khác biệt giữa dữ liệu được khai báo trực tiếp và dữ liệu nhận qua kế thừa. Các Condition hiển thị riêng, giúp tránh hiểu sai dấu hiệu nhận biết thành tính chất được kế thừa.')
    h(d,'7.2 Truy vấn và giải thích bằng đường đi',2,page=True)
    figure(d,'12_query_Q15_graph.jpg','Hình 4. Q15 chạy thật và hiển thị hai đường đi giải thích tính chất.')
    para(d,'Trang Truy vấn trình bày câu Cypher, tham số, bảng kết quả và graph khi có dữ liệu graph. Kết quả Q15 giữ đường đi nên người dùng có thể kiểm tra từng quan hệ, thay vì chỉ nhận một tên tính chất.')
    h(d,'7.3 Đối chiếu trực tiếp trên Neo4j',2,page=True)
    figure(d,'14_neo4j_query.jpg','Hình 5. Neo4j Query/Browser trả HTHOI và HV cho điều kiện bốn cạnh bằng nhau.',width=8.5,max_height=17)
    para(d,'Ảnh Browser dùng truy vấn đọc có traversal tương đương Q06 trên dataset thật. Kết quả khớp với API. Validation và kết quả dài được lưu trong báo cáo JSON để có thể đối chiếu độc lập với giao diện.')
    h(d,'CHƯƠNG 8 — KIỂM THỬ',page=True)
    table(d,'Bảng 7. Kết quả kiểm thử đã lưu trong project',['Giai đoạn','Nội dung','Kết quả'],[
        ('Phase 2','20 validation; 15 truy vấn Cypher','20/20; 15/15 PASS'),
        ('Phase 3','pytest: 69 integration + 22 unit','91/91 PASS'),
        ('Phase 3','HTTP smoke','38/38 PASS'),
        ('Phase 5','pytest: 69 integration + 23 unit','92/92 PASS'),
        ('Phase 5','HTTP chính; HTTP tìm kiếm/lỗi','41/41; 8/8 PASS'),
        ('Phase 5','Kiểm thử giao diện','78/78 PASS'),
        ('Phase 5','Môi trường ảo mới, dependencies khai báo','92 pytest; 41 HTTP PASS'),
        ('Phase 5','Seed lại hai lần; so snapshot','0 node/0 cạnh mới; không đổi dữ liệu')],[2,7.5,7])
    para(d,'Validation kiểm tra ID trùng, node thiếu quan hệ phù hợp, self-loop, chu trình IS_A, cạnh sai loại và số liệu mong đợi. Cả 20 phép kiểm tra đều có violations = 0. Việc đọc schema xác nhận ba unique constraint và bốn index ONLINE, bao gồm index hỗ trợ constraint.')
    para(d,'System test kiểm tra kết nối thật, truy vấn, route, tìm kiếm, graph và các tình huống Neo4j lỗi. Khi Neo4j không khả dụng, API trả 503 và giao diện báo lỗi kết nối. Kiểm thử không thay Neo4j bằng database khác và không giả lập kết quả toán học. Môi trường ảo mới không đồng nghĩa đã kiểm thử trên một máy hoặc database hoàn toàn mới.')
    para(d,'Nguồn minh chứng: docs/PHASE_2_RESULTS.json; docs/PHASE_3_RESULTS.json; docs/PHASE_5_RESULTS.json; các báo cáo HTTP, UI và XML pytest tương ứng. Ảnh chụp chứng minh giao diện; số lượng test dựa trên các báo cáo kiểm thử, không suy từ ảnh.')
    h(d,'CHƯƠNG 9 — ĐÁNH GIÁ',page=True)
    for title,text in [
        ('9.1 Kết quả đạt được','Project đã kết nối Neo4j thật, seed idempotent và cung cấp bảy loại tứ giác cùng 19 Property, 18 Condition. Website chạy các chức năng đã thiết kế: tổng quan, danh sách, chi tiết, tìm kiếm, taxonomy, lân cận và 15 truy vấn demo. Mô hình thể hiện đa kế thừa, dùng chung Property và điều kiện có ngữ cảnh.'),
        ('9.2 Hạn chế','Phạm vi chỉ có bảy loại tứ giác và tri thức được nhập có chủ đích. Hệ thống chưa là theorem prover, chưa xác nhận tính chất từ toạ độ hoặc ảnh, chưa suy luận mọi hệ quả từ các giả thiết. Quy ước hình thang của project cần được trình bày trước khi demo. Server Flask hiện phục vụ chạy cục bộ, chưa phải cấu hình triển khai công khai.'),
        ('9.3 Hướng phát triển','Có thể bổ sung các loại hình khác, ghi nguồn cho từng mệnh đề, hỗ trợ tra cứu điều kiện theo tập giả thiết và kiểm tra bằng dữ liệu toạ độ. Khi mở rộng, cần kiểm tra tính nhất quán giữa các quy ước phân loại. Việc triển khai nhiều người dùng hoặc công bố GitHub cần các bước vận hành và audit riêng.')]:
        h(d,title,2); para(d,text)
    h(d,'KẾT LUẬN')
    para(d,'Đề tài dùng Neo4j để mô hình hoá kiến thức hình học theo các quan hệ có thể truy vấn. Hai đường kế thừa của Hình vuông và truy vấn giải thích tính chất cho thấy đồ thị mang thông tin thực sự, thay vì chỉ dùng node để thay thế bảng. Kết quả kiểm thử xác nhận phiên bản hiện tại phù hợp cho demo môn NoSQL. Hai tài liệu Word và source được chuẩn bị cho bài nộp; link GitHub còn chờ Phase 7.')
    h(d,'TÀI LIỆU THAM KHẢO',page=True)
    refs = [
        ('Neo4j — Cypher patterns','https://neo4j.com/docs/cypher-manual/current/patterns/'),
        ('Neo4j — Create constraints','https://neo4j.com/docs/cypher-manual/current/schema/constraints/create-constraints/'),
        ('Neo4j — Python Driver, Execute queries','https://neo4j.com/docs/python-manual/current/query-simple/'),
        ('Neo4j Desktop — Installation','https://neo4j.com/docs/desktop/current/installation/'),
        ('vis-network — Layout','https://visjs.github.io/vis-network/docs/network/layout.html')]
    for shape in SHAPES:
        if shape['source_ref'].startswith('https:'): refs.append(('Wolfram MathWorld — '+shape['name'],shape['source_ref']))
    refs.append(('Wolfram MathWorld — Trapezoid','https://mathworld.wolfram.com/Trapezoid.html'))
    for i,(title,url) in enumerate(refs,1):
        para(d,f'[{i}] {title}.'); hyperlink(d,url,url)
    para(d,'Nguồn nội bộ: PHASE_1_PHAN_TICH.md; database/shapes.json, properties.json, conditions.json; database/demo_queries.cypher; báo cáo Phase 2, 3, 5. Các định nghĩa và dấu hiệu nhận biết trong báo cáo theo đúng catalog và quy ước đã nghiệm thu của project.')
    d.save(DOCS/'TONG_QUAN_DU_AN.docx')


def hyperlink(d,label,url):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    link = OxmlElement('w:hyperlink')
    rid = p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link.set(qn('r:id'),rid)
    run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'),'20'); pr.append(size)
    run.append(pr); text = OxmlElement('w:t'); text.text = label; run.append(text); link.append(run); p._p.append(link)


def manual():
    d = new_doc('HƯỚNG DẪN SỬ DỤNG')
    h(d,'1. Giới thiệu',page=True)
    para(d,'Tài liệu hướng dẫn chạy và demo project trên Windows PowerShell. Neo4j là database chính. Các bước cài đặt dùng dữ liệu và script có trong source; ảnh là kết quả thật của phiên bản đã kiểm thử. Trước khi nộp, điền các thông tin cá nhân trên trang bìa. Link GitHub sẽ được bổ sung sau audit Phase 7.')
    h(d,'2. Yêu cầu môi trường')
    table(d,'Bảng 1. Môi trường đã kiểm thử',['Thành phần','Phiên bản / yêu cầu'], [('Hệ điều hành','Windows; PowerShell'),('Python','3.13.5 đã kiểm thử'),('Neo4j','Enterprise 2026.09.0; database quadrilateral'),('Thư viện Python','Theo requirements.txt'),('Trình duyệt','Trình duyệt hiện đại có JavaScript'),('Kết nối','Neo4j Bolt tại localhost:7687; web Flask :5000')],[4,12.5])
    para(d,'Internet cần cho cài đặt công cụ và thư viện. Demo sử dụng Neo4j cục bộ; vis-network đã có trong static. Các phiên bản hoặc hệ điều hành khác chưa được xác nhận bằng báo cáo hiện tại.')
    h(d,'3. Cấu trúc thư mục')
    code(d,'nosql/\n  app/                 # backend, templates, static\n  database/            # JSON, schema, seed, query, validation\n  docs/                # Word, báo cáo, ảnh thật\n  scripts/             # seed và kiểm thử\n  tests/               # unit và integration\n  .env.example\n  requirements.txt\n  README.md\n  run.py')
    para(d,'Các lệnh dưới đây chạy từ thư mục gốc project. Không đưa .env, .venv hoặc thư mục .qa lên GitHub. Thư mục .qa chứa dữ liệu kiểm tra cục bộ và đã được loại trừ trong .gitignore.')
    h(d,'4. Khởi động Neo4j',page=True)
    para(d,'Cài Neo4j Desktop từ nguồn chính thức. Tạo instance cục bộ, đặt tên tuỳ chọn và tự đặt mật khẩu cho tài khoản neo4j. Chọn phiên bản Enterprise có hỗ trợ nhiều database phù hợp giấy phép phát triển. Bấm Start và đợi trạng thái Running trước khi chạy script. Tên instance không phải tên database.')
    hyperlink(d,'Hướng dẫn cài Neo4j Desktop','https://neo4j.com/docs/desktop/current/installation/')
    h(d,'5. Tạo/chọn database quadrilateral')
    para(d,'Trong công cụ Query/Browser của instance đang chạy, chọn system bằng bộ chọn database rồi tạo database riêng. Không xóa hoặc dùng lại dữ liệu của database khác. Người dùng cần quyền tạo database. Phiên bản đã kiểm thử là Enterprise; không mặc định Community hỗ trợ cấu hình nhiều database này.')
    code(d,'CREATE DATABASE quadrilateral IF NOT EXISTS;\nSHOW DATABASES;')
    para(d,'Sau khi database quadrilateral online, chọn quadrilateral trước khi chạy query dữ liệu. Bộ chọn database phải hiển thị đúng tên. Một instance có thể chứa system, neo4j và quadrilateral cùng lúc.')
    h(d,'6. Cấu hình .env')
    code(d,'Copy-Item .env.example .env\n\nNEO4J_URI=bolt://localhost:7687\nNEO4J_USER=neo4j\nNEO4J_PASSWORD=your_password\nNEO4J_DATABASE=quadrilateral')
    para(d,'Thay your_password bằng mật khẩu do bạn đặt; không lưu mật khẩu thật vào README hoặc tài liệu. .env nằm ở thư mục gốc, không phải .env.txt. URI không chứa username/password. Flask ưu tiên biến môi trường đã có, còn CLI Phase 2 đọc cấu hình từ .env; nếu kết nối khác nhau, kiểm tra biến môi trường trong terminal. Không chọn system làm database ứng dụng.')
    h(d,'7. Cài dependencies',page=True)
    code(d,'py -3.13 -m venv .venv\n'+PY+' -m pip install -r requirements.txt\n'+PY+' -m pip check')
    para(d,'Nếu máy chỉ có lệnh python, dùng python -m venv .venv. Những lệnh tiếp theo gọi trực tiếp Python trong môi trường ảo nên không cần kích hoạt môi trường. requirements.txt gồm neo4j, python-dotenv, Flask và pytest với phiên bản cố định.')
    h(d,'8. Seed dữ liệu')
    para(d,'Thực hiện theo thứ tự: kiểm tra kết nối, tạo constraint/index, sinh seed từ catalog rồi seed bằng MERGE. Không sửa ID của node đã có nếu chưa kiểm tra tác động lên relationship.')
    code(d,'\n'.join(PY+' scripts/phase2.py '+arg for arg in ['connect','schema','build-seed','seed','finish']))
    para(d,'Kết quả đúng của dataset hiện tại là 7 Shape, 19 Property, 18 Condition: tổng 44 node và 79 relationship. finish kiểm tra, seed lần hai và đối chiếu idempotent rồi lưu báo cáo. Seed lại không tạo dữ liệu trùng. Nếu lệnh báo lỗi, đọc và xử lý nguyên nhân trước khi tiếp tục; không thay database bằng dữ liệu giả để demo.')
    h(d,'9. Kiểm tra validation')
    code(d,PY+' scripts/phase2.py verify')
    para(d,'Lệnh kiểm tra 20 validation và 15 query, đối chiếu expected_results.json. Báo cáo nghiệm thu hiện tại là 20/20 validation, 15/15 query PASS. Kết quả verify mới phản ánh môi trường tại thời điểm chạy; nên giữ lại các báo cáo lịch sử trước khi ghi đè. Có thể kiểm tra database/validation.cypher trong Neo4j Browser khi đã khai báo tham số.')
    h(d,'10. Chạy Flask',page=True)
    code(d,PY+' run.py')
    para(d,'Giữ cửa sổ terminal này mở. Server chạy tại 127.0.0.1:5000, debug và reloader tắt. Nếu port 5000 đang dùng, chọn port khác và truy cập đúng port đó:')
    code(d,PY+' run.py --port 5001')
    h(d,'11. Truy cập website')
    hyperlink(d,'http://127.0.0.1:5000','http://127.0.0.1:5000')
    para(d,'Mở địa chỉ trên bằng trình duyệt. Navbar dẫn tới Tổng quan, Tứ giác, Graph, Truy vấn và Giới thiệu. Kiểm tra trạng thái kết nối Neo4j và số liệu dashboard trước khi bắt đầu thuyết trình.')
    para(d,'Luồng demo gợi ý: Tổng quan → taxonomy → tìm hinh vuong → chi tiết Hình vuông → Property kế thừa → Condition → Q06, Q09, Q15 → đối chiếu Neo4j Browser. Dùng graph để giải thích quan hệ, không chỉ đọc các card thống kê.')
    ui = [
        (12,'Trang Tổng quan','Mở /. Đọc các card node, relationship và số loại hình; kết quả đúng là 44, 79 và 7. Xem graph tổng quan rồi chọn node để đọc thông tin. Các thống kê chỉ thuộc dataset quadrilateral-v1.','01_dashboard.jpg','Hình 1. Trang Tổng quan lấy số liệu từ Neo4j thật.'),
        (13,'Danh sách tứ giác','Mở /shapes. Trang hiển thị bảy loại hình với định nghĩa và lối mở chi tiết. Đối chiếu tên Hình thang và Hình bình hành với hai nhánh riêng của taxonomy.','03_shapes.jpg','Hình 2. Danh sách bảy loại tứ giác.'),
        (14,'Tìm kiếm','Trên /shapes, nhập hinh vuong vào ô tìm kiếm. Kết quả là Hình vuông (HV). Thử Hình vuông để kiểm tra tìm kiếm có dấu; xoá nội dung để trở lại bảy hình. Không có kết quả thì trang thông báo rõ, không giữ kết quả của lần tìm trước.',IMAGES/'04_search_hinh_vuong.jpg','Hình 3. Tìm kiếm không dấu hinh vuong trả đúng Hình vuông.'),
        (15,'Chi tiết Hình vuông','Chọn Hình vuông hoặc mở /shapes/HV. Đọc định nghĩa, hai lớp cha Hình chữ nhật và Hình thoi, rồi xem các tổ tiên. HV có HBH, HCN, HTHOI, TQ là tổ tiên; không có HT hoặc HTC.','04_shape_hv_overview.jpg','Hình 4. Trang chi tiết Hình vuông và quan hệ phân loại.'),
        (16,'Xem Property kế thừa','Cuộn tới phần tính chất. HV có 12 tính chất hiệu lực dù không có HAS_PROPERTY trực tiếp. Đọc nguồn khai báo ở các lớp cha/tổ tiên. Đây là phần cần giải thích khi giảng viên hỏi vì sao dùng graph.','05_shape_hv_properties.jpg','Hình 5. Tính chất kế thừa và nguồn khai báo của Hình vuông.'),
        (17,'Xem Condition','Đọc năm cách nhận biết Hình vuông. Trong mỗi Condition, ngữ cảnh và các giả thiết nối bằng AND; năm Condition là các lựa chọn OR. Ví dụ C_HV_05 cần Tứ giác, bốn cạnh bằng nhau và bốn góc vuông đồng thời. Ứng dụng trình bày tri thức, chưa tự phân loại hình từ đầu vào.','06_shape_hv_conditions.jpg','Hình 6. Các Condition nhận biết Hình vuông.'),
        (18,'Graph taxonomy','Mở /graph và chọn chế độ phân loại. Kiểm tra 7 Shape, 7 cạnh IS_A. Nhìn mũi tên HV tới HCN và HTHOI; hai nhánh cùng tới HBH rồi TQ. Kéo hoặc thu phóng để quan sát, chọn node/cạnh để xem chi tiết.','02_taxonomy_graph.jpg','Hình 7. Đồ thị taxonomy với đa kế thừa.'),
        (19,'Graph neighborhood','Chọn graph lân cận, chọn Hình vuông và độ sâu 1. Kết quả là 8 node/11 cạnh. Tăng lên 2 để được 23/43, và lên 3 để được 39/71. Node gồm Shape, Property, Condition; graph lấy các cạnh giữa node được chọn, không chỉ các cạnh trên một đường duyệt.','07_graph_depth1.jpg','Hình 8. Graph lân cận HV ở độ sâu 1.'),
        (20,'Truy vấn demo','Mở /queries. Chọn query, xem Cypher và tham số rồi thực thi. Q06 trả Hình thoi và Hình vuông; Q09 trả hai đường đi TQ–HV; Q15 giải thích một tính chất qua hai nhánh kế thừa. Chỉ có query trong danh mục được chạy, không có ô thực thi Cypher tuỳ ý.','10_query_Q06.jpg','Hình 10. Q06 chạy thật: các hình có bốn cạnh bằng nhau.'),
        (21,'Neo4j Query/Browser','Mở Query/Browser của instance trong Neo4j Desktop. Chọn database quadrilateral. Khai báo tham số bằng lệnh ở mục 22, sau đó copy query từ database/demo_queries.cypher. Dùng Table để đối chiếu giá trị hoặc Graph khi kết quả có node/relationship/path.','14_neo4j_query.jpg','Hình 12. Đối chiếu truy vấn bốn cạnh bằng nhau trực tiếp trên Neo4j.')]
    for num,title,text,img,caption in ui:
        h(d,f'{num}. {title}',page=True); para(d,text)
        portrait = num in [14,21]
        figure(d,img,caption,width=8.5 if portrait else 16,max_height=17 if portrait else 15.5)
        if num == 19:
            h(d,'19.1 Tăng độ sâu và đọc quan hệ',2,page=True)
            figure(d,'08_graph_depth2.jpg','Hình 9. HV ở độ sâu 2: 23 node và 43 relationship.')
            para(d,'Số node/cạnh tăng do graph mở rộng qua các loại relationship theo cả hai hướng. Chọn từng cạnh để đọc nhãn; không suy ra quan hệ IS_A chỉ vì hai node gần nhau trong bố cục. Nút căn graph đưa tất cả node vào khung xem.')
        if num == 20:
            h(d,'20.1 Đọc đường đi giải thích Q15',2,page=True)
            figure(d,'12_query_Q15_graph.jpg','Hình 11. Hai đường đi của Q15 trên ứng dụng.')
            para(d,'Đọc mỗi đường: HV → HCN/HTHOI → HBH → Property Đường chéo chia đôi nhau. Các cạnh trước là IS_A; cạnh cuối là HAS_PROPERTY. Kết quả trùng Property được giữ ở dạng đường đi vì hai đường là hai chứng cứ về kế thừa.')
        if num == 21:
            h(d,'21.1 Xem graph trên Neo4j',2,page=True)
            figure(d,'13_neo4j_graph.jpg','Hình 13. Graph taxonomy thật trong Neo4j Query/Browser.',width=8.5,max_height=17)
            para(d,'Ảnh Browser chụp ở cửa sổ hẹp nên nhãn có thể nhỏ. Dùng graph taxonomy của ứng dụng để đọc tên rõ hơn; Browser dùng để chứng minh dữ liệu node/relationship được truy vấn trực tiếp từ Neo4j.')
    h(d,'22. Một số Cypher mẫu',page=True)
    para(d,'Khai báo trong Neo4j Query/Browser trước khi chạy (không đưa lệnh :params vào Neo4j Driver):')
    code(d,PARAM)
    for q in ['Q03','Q06','Q09']:
        h(d,q+' — '+{a:b for a,b,c in QUERY_INFO}[q],2)
        code(d,QUERIES[q]); para(d,'Kết quả đã kiểm thử: '+{a:c for a,b,c in QUERY_INFO}[q]+'.')
    para(d,'Q02 có 12 tính chất; Q08 trả HCN, HTC, HV; Q12 trả HTHOI, HV; Q15 có hai đường giải thích. Toàn bộ 15 query và chú thích ở database/demo_queries.cypher, kết quả chuẩn ở database/expected_results.json.')
    h(d,'23. Dừng ứng dụng',page=True)
    para(d,'Trong terminal đang chạy Flask, nhấn Ctrl+C và đợi server thoát. Đóng tab web khi không cần demo. Nếu muốn dừng Neo4j, dùng Stop trong Neo4j Desktop sau khi các lệnh đang chạy đã kết thúc. Không chọn Delete instance/database để dừng dịch vụ.')
    h(d,'24. Xử lý lỗi thường gặp')
    table(d,'Bảng 2. Kiểm tra và xử lý lỗi',['Hiện tượng','Cách xử lý'],[
        ('Neo4j không kết nối / HTTP 503','Kiểm tra instance Running. Thử scripts/phase2.py connect; đọc lỗi đã được che thông tin nhạy cảm.'),
        ('Sai mật khẩu','Kiểm tra username/password trong .env, dùng mật khẩu đã đặt cho instance; không gửi mật khẩu vào báo cáo hoặc ảnh.'),
        ('Database không tồn tại','Kiểm tra SHOW DATABASES trong system; tạo quadrilateral với quyền phù hợp rồi chọn đúng database.'),
        ('Không truy cập Bolt port 7687','Kiểm tra địa chỉ Bolt của instance có khớp NEO4J_URI; instance và firewall phải cho phép kết nối cục bộ.'),
        ('Thiếu .env / cấu hình khác giữa CLI và web','Đặt .env đúng thư mục gốc, kiểm tra đuôi .txt và biến môi trường có sẵn trong terminal.'),
        ('Thiếu package / ModuleNotFoundError','Dùng đúng Python .venv để cài requirements.txt rồi pip check; không cài vào Python khác.'),
        ('Port 5000 bận','Dừng server cũ nếu chính bạn đang dùng hoặc chạy run.py --port 5001; mở URL đúng port.'),
        ('Query có lỗi thiếu tham số','Chọn quadrilateral và chạy :params ở mục 22 trước khi copy Cypher trong Browser.')],[5,11.5])
    para(d,'Sau khi sửa cấu hình, khởi động lại Flask để nạp cấu hình mới. Không bỏ qua lỗi validation hoặc thay expected result để làm test PASS. Nếu chưa rõ nguyên nhân, giữ nguyên dữ liệu và xem báo cáo kiểm thử/README trước khi sửa schema.')
    d.save(DOCS/'HUONG_DAN_SU_DUNG.docx')


def finalize_repository_links():
    """Update the approved report with the repository published in Phase 7."""
    replacements = {
        'Việc công bố link GitHub sẽ thực hiện sau kiểm tra secret ở Phase 7.':
            'Data + Source App được lưu tại GitHub: https://github.com/phi1411/neo4j.',
        'Việc triển khai nhiều người dùng hoặc công bố GitHub cần các bước vận hành và audit riêng.':
            'Việc triển khai nhiều người dùng cần các bước vận hành và audit riêng.',
        'Hai tài liệu Word và source được chuẩn bị cho bài nộp; link GitHub còn chờ Phase 7.':
            'Hai tài liệu Word và source được chuẩn bị cho bài nộp. Repository: https://github.com/phi1411/neo4j.',
        'Link GitHub sẽ được bổ sung sau audit Phase 7.':
            'Repository: https://github.com/phi1411/neo4j.',
    }
    for name in ['TONG_QUAN_DU_AN.docx','HUONG_DAN_SU_DUNG.docx']:
        path = DOCS/name
        d = Document(path)
        for p in d.paragraphs:
            if any(old in p.text for old in replacements):
                text = p.text
                for old,new in replacements.items(): text = text.replace(old,new)
                p.text = text
        d.save(path)


if __name__ == '__main__':
    overview()
    manual()
    finalize_repository_links()
    print(json.dumps({'created':['docs/TONG_QUAN_DU_AN.docx','docs/HUONG_DAN_SU_DUNG.docx'],'nodes':44,'relationships':79},ensure_ascii=False))
